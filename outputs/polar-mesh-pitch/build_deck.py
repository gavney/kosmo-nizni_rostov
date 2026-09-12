"""Build Polar Mesh 5-minute pitch deck."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screens"
ASSETS = Path(r"C:\Users\lebed\.cursor\projects\c-Users-lebed-Desktop-kosmoshack2026\assets")
OUT = ROOT / "Polar-Mesh-zashchita-5min.pptx"

W, H = Inches(13.333), Inches(7.5)
DARK = RGBColor(0x0B, 0x0F, 0x16)
INK = RGBColor(0xF4, 0xF6, 0xFA)
MUTED = RGBColor(0xA8, 0xB3, 0xC4)
ACCENT = RGBColor(0x6E, 0xC8, 0xFF)
WARM = RGBColor(0xFF, 0xD4, 0x5A)
LIGHT_BG = RGBColor(0xF3, 0xF5, 0xF8)
LIGHT_INK = RGBColor(0x14, 0x1A, 0x24)
LIGHT_MUTED = RGBColor(0x5A, 0x66, 0x7A)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(shape, lines, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color, font=font)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_image(slide, path: Path, left, top, width, height):
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def card(slide, left, top, width, height, fill=RGBColor(0x14, 0x1A, 0x26)):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    add_bg(s, DARK)
    add_image(s, SHOTS / "pitch-06-globe-hero.png", Inches(4.8), 0, Inches(8.6), H)
    shade = card(s, 0, 0, Inches(6.4), H, RGBColor(0x0B, 0x0F, 0x16))
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(5.0), Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "POLAR MESH"
    set_run(r, 14, True, ACCENT)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Связь для Севера,\nкоторую можно\nспроектировать"
    set_run(r, 34, True, INK)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "\nКоманда Ростов · КосмоХакатон 2026"
    set_run(r, 16, False, MUTED)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Цель: сквозная доступность ≥ 90%"
    set_run(r, 16, False, WARM)

    # 2 Task
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT_BG)
    add_image(s, ASSETS / "deck-bg-light.png", 0, 0, W, H)
    title = s.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(7), Inches(1))
    add_text(title, "О чём задача", size=32, bold=True, color=LIGHT_INK)
    sub = s.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(6.2), Inches(1.2))
    add_text(
        sub,
        "Северным клиентам нужен не «спутник в небе», а путь до шлюза через группировку. Видеть аппарат — ещё не значит доставить данные.",
        size=18,
        color=LIGHT_MUTED,
    )
    bullets = [
        ("C65 · C70 · C72", "северные терминалы"),
        ("ISL-меш", "ретрансляция только через спутники"),
        ("G_MUR", "выход в наземную сеть"),
    ]
    for i, (h, t) in enumerate(bullets):
        c = card(s, Inches(0.7 + i * 2.7), Inches(3.0), Inches(2.5), Inches(2.2), RGBColor(0xFF, 0xFF, 0xFF))
        c.shadow.inherit = False
        tb = s.shapes.add_textbox(Inches(0.9 + i * 2.7), Inches(3.25), Inches(2.2), Inches(1.8))
        add_text(tb, [h, "", t], size=16, bold=False, color=LIGHT_INK)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    add_image(s, SHOTS / "pitch-06-globe-hero.png", Inches(8.5), Inches(1.5), Inches(4.4), Inches(4.8))

    # 3 Demo UI
    s = prs.slides.add_slide(blank)
    add_bg(s, DARK)
    add_image(s, SHOTS / "pitch-01-orbit.png", Inches(0.4), Inches(1.35), Inches(12.5), Inches(5.7))
    bar = card(s, 0, 0, W, Inches(1.15), RGBColor(0x0B, 0x0F, 0x16))
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    add_text(t, "Интерфейс, в котором можно пройти весь сценарий жюри", size=26, bold=True, color=INK)
    note = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12), Inches(0.4))
    add_text(note, "Глобус · таймлайн · метрики доступности · маршрут в один взгляд", size=14, color=MUTED)

    # 4 Design
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT_BG)
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(6), Inches(0.8))
    add_text(t, "Проектируем, а не просто смотрим", size=28, bold=True, color=LIGHT_INK)
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(5.8), Inches(2.2))
    add_text(
        body,
        [
            "Меняем очередь запуска и ориентацию плоскостей.",
            "Считаем сутки на сетке кейса.",
            "Сразу видим, тянем ли цель 90% — и где рвётся связь.",
        ],
        size=18,
        color=LIGHT_MUTED,
    )
    metrics = s.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(5.8), Inches(2.5))
    add_text(
        metrics,
        [
            "На полной группировке сейчас:",
            "C65 — 96.7%   C70 — 98.8%   C72 — 98.9%",
            "Цель кейса — 90%. Мы её перекрываем.",
        ],
        size=18,
        color=LIGHT_INK,
    )
    add_image(s, SHOTS / "pitch-03-project.png", Inches(6.6), Inches(1.0), Inches(6.3), Inches(5.6))

    # 5 Sats / chain
    s = prs.slides.add_slide(blank)
    add_bg(s, DARK)
    add_image(s, SHOTS / "pitch-02-sats.png", Inches(0.35), Inches(1.2), Inches(12.6), Inches(5.85))
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    add_text(t, "Кто реально несёт трафик — видно сразу", size=26, bold=True, color=INK)
    note = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12), Inches(0.4))
    add_text(note, "Жёлтым — аппараты в текущей цепи. Остальные спокойнее, но читаемы даже на ночной стороне.", size=14, color=MUTED)

    # 6 Resilience
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT_BG)
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.7))
    add_text(t, "Ломаем сеть специально — и смотрим, что остаётся", size=28, bold=True, color=LIGHT_INK)
    add_image(s, SHOTS / "pitch-04-ground.png", Inches(0.5), Inches(1.3), Inches(7.8), Inches(5.5))
    side = s.shapes.add_textbox(Inches(8.5), Inches(1.6), Inches(4.3), Inches(4.5))
    add_text(
        side,
        [
            "Отказ КА",
            "Отключение шлюза",
            "Сигнал «нет связи»",
            "",
            "Показываем не только «упало»,",
            "но и почему: нет видимости,",
            "дыра в ISL или шлюз offline.",
        ],
        size=18,
        color=LIGHT_INK,
    )

    # 7 Compare
    s = prs.slides.add_slide(blank)
    add_bg(s, DARK)
    add_image(s, SHOTS / "pitch-05-compare.png", Inches(0.35), Inches(1.2), Inches(12.6), Inches(5.85))
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    add_text(t, "Два варианта рядом — выбираем глазами и цифрами", size=26, bold=True, color=INK)
    note = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12), Inches(0.4))
    add_text(note, "Сохранили конфиги → сравнили availability и перерывы → оставили тот, что держит север.", size=14, color=MUTED)

    # 8 Tech
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT_BG)
    t = s.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(0.8))
    add_text(t, "Под капотом — без магии, по правилам кейса", size=28, bold=True, color=LIGHT_INK)
    items = [
        ("Геометрия", "официальный geometry.py организаторов"),
        ("Маршруты", "BFS по hops и Dijkstra по километрам"),
        ("Правило сети", "земля не ретранслирует — только КА"),
        ("Выгрузка", "cosmo-A-result-1.0 + сценарий целиком"),
    ]
    for i, (h, b) in enumerate(items):
        y = Inches(1.6 + i * 1.25)
        c = card(s, Inches(0.7), y, Inches(12), Inches(1.1), RGBColor(0xFF, 0xFF, 0xFF))
        tb = s.shapes.add_textbox(Inches(1.0), y + Inches(0.22), Inches(11.4), Inches(0.8))
        add_text(tb, [h, b], size=16, color=LIGHT_INK)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1A, 0x6F, 0xB5)

    # 9 Close
    s = prs.slides.add_slide(blank)
    add_bg(s, DARK)
    add_image(s, SHOTS / "pitch-06-globe-hero.png", 0, 0, W, H)
    overlay = card(s, Inches(1.8), Inches(2.0), Inches(9.7), Inches(3.5), RGBColor(0x0B, 0x0F, 0x16))
    tb = s.shapes.add_textbox(Inches(2.2), Inches(2.35), Inches(8.9), Inches(3.0))
    add_text(
        tb,
        [
            "Polar Mesh",
            "",
            "Спроектировать · проверить на отказах ·",
            "сравнить · выгрузить результат.",
            "",
            "Команда Ростов",
        ],
        size=22,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(34)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
